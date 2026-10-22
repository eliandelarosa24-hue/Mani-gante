# Genera el guion de exposición y la matriz de cumplimiento (Word)
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_DIR = sys.argv[1]
VIOLET = RGBColor(0x6D, 0x28, 0xD9)
BLUE = RGBColor(0x1D, 0x4E, 0xD8)
GREY = RGBColor(0x55, 0x5B, 0x75)
AMBER = RGBColor(0xB4, 0x53, 0x09)


def base_doc():
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.2))
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.paragraph_format.space_after = Pt(4)
    for name, size, color in (("Heading 1", 18, VIOLET), ("Heading 2", 14, BLUE), ("Heading 3", 12, VIOLET)):
        h = d.styles[name]
        h.font.name = "Arial"
        h.font.size = Pt(size)
        h.font.color.rgb = color
        h.font.bold = True
        h.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    return d


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def table(d, header, rows, widths=None, size=9.5, status_col=None):
    t = d.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, "4C1D95")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(size)
            if status_col is not None and i == status_col:
                r.bold = True
                r.font.color.rgb = {"Cumple": RGBColor(0x04, 0x78, 0x57), "Parcial": AMBER}.get(str(v).split(" ")[0], AMBER)
    if widths:
        t.autofit = False
        grid = t._tbl.tblGrid
        for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
            gc.set(qn("w:w"), str(int(widths[i] * 567)))
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    # repetir encabezado
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)
    return t


def para(d, text, bold_lead=None, color=None, italic=False, size=None):
    p = d.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
        if color:
            r.font.color.rgb = color
    r = p.add_run(text)
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    return p


def bullet(d, text, lead=None):
    p = d.add_paragraph(style="List Bullet")
    if lead:
        r = p.add_run(lead)
        r.bold = True
    p.add_run(text)
    return p


# ---------------------------------------------------------------------------
# Guion
# ---------------------------------------------------------------------------
SLIDES = [
    (1, "Portada", 1, 20,
     "Buenos días, profesora y compañeros. Somos el equipo de Sound Travel. Durante esta unidad investigamos cómo coordinar un servicio turístico de eventos en nuestra región y diseñamos esta agencia, que es ficticia: elegimos la Opción 2 de la guía, que pedía crear una empresa contextualizada en nuestro territorio.",
     ["Empresa ficticia, Opción 2.", "Trabajo de investigación propio del equipo."],
     "Partamos por el problema que encontramos."),
    (2, "El problema y nuestra respuesta: Sound Travel", 1, 60,
     "Al investigar el contexto vimos que el turismo en Chile está creciendo: según la Subsecretaría de Turismo, en 2023 llegaron más de 3,6 millones de visitantes internacionales, y además hay mucha gente que viaja entre regiones para ver conciertos. "
     "Pero en Temuco, cada vez que hay un evento masivo en el Germán Becker o en la Isla Cautín, se repiten tres problemas: la micro deja de pasar temprano, los estacionamientos colapsan y las aplicaciones suben sus tarifas. "
     "Nuestra respuesta es Sound Travel: una agencia de turismo de eventos con base en Temuco, que opera rutas a Padre Las Casas, Lautaro, Villarrica y Pucón, y viajes especiales a Concepción y Santiago. "
     "Tomamos como escenario el Festival Internacional de Temuco 2027, pero aclaramos que no encontramos confirmación oficial; lo usamos solo para planificar.",
     ["Tres problemas concretos del territorio.", "Localización: Temuco como base.", "El festival 2027 es un escenario, no un hecho."],
     "Ahora, ¿qué busca la empresa y para quién trabaja?"),
    (3, "Identidad y propuesta de valor", 1, 50,
     "Nuestra misión es ofrecer traslado y actividades complementarias para asistentes a eventos, garantizando seguridad, comodidad y un viaje bien planificado. Y queremos ser reconocidos en el sur de Chile justamente por la buena organización. "
     "Definimos dos públicos: jóvenes de 18 a 35 años que van a conciertos y festivales, y grupos o parejas de hasta 45 años que buscan panoramas de fin de semana, como un sunset en la zona lacustre. "
     "A ambos les ofrecemos el Pack Sound Travel, que junta cuatro cosas: traslado ida y vuelta, entrada asegurada con ticketeras oficiales, un guía que los acompaña y paradas gastronómicas. En una frase: llegar puntual, disfrutar tranquilo y volver seguro.",
     ["Misión y perfiles apuntan a seguridad y comodidad.", "Cuatro componentes en un solo pack."],
     "Para cumplir esta promesa necesitamos una organización clara. Le paso la palabra al Integrante 2."),
    (4, "Estructura organizacional y recurso humano", 2, 45,
     "Gracias. Diseñamos una estructura simple y vertical. Arriba está la Dirección General, que planifica, busca alianzas y vela por la normativa. Debajo hay tres áreas: Operaciones y Logística, que arma rutas y revisa los minibuses; Comercial, que atiende redes y reservas; y Administración y Finanzas, que lleva la contabilidad y paga a los proveedores. "
     "Decidimos que el Equipo de Terreno, guías y conductores, dependa de Operaciones, porque así quien planifica la ruta coordina directamente a quienes están con los pasajeros.",
     ["Cinco cargos con funciones claras.", "Terreno depende de Operaciones."],
     "Esta misma estructura la podemos mirar por funciones."),
    (5, "Organigrama funcional y procesos internos y externos", 2, 40,
     "En el organigrama funcional agrupamos el trabajo en cuatro funciones: directiva, operativa, comercial y administrativa, y todas apuntan al mismo servicio integrado. "
     "Estas funciones se traducen en procesos internos, como planificar rutas, gestionar reservas o controlar pagos, y en procesos externos en tres frentes: proveedores, alianzas con instituciones y clientes. Lo interno y lo externo se retroalimentan: lo que preparamos adentro es lo que comprometemos afuera.",
     ["Jerárquico = quién depende de quién; funcional = qué hace cada área.", "Tres frentes externos."],
     "Todo esto se ordena en un flujo de cinco etapas."),
    (6, "Diagrama de flujo del servicio en cinco etapas", 2, 45,
     "Este es el recorrido completo. Primero difundimos el Pack Evento en redes; luego el cliente cotiza y compra por web o WhatsApp; en la etapa 3 bloqueamos asientos, comidas y entradas con los proveedores; en la 4 se hace el traslado con el guía; y en la 5 volvemos seguros a Temuco y enviamos una encuesta. "
     "Esa encuesta vuelve al inicio, por eso la flecha verde: es mejora continua. Al revisar el informe notamos que no decía quién envía la encuesta; lo dejamos marcado con asterisco.",
     ["La etapa 3 concentra la coordinación externa.", "La encuesta cierra el ciclo."],
     "En la etapa 3 aparecen muchos actores. Le paso la palabra al Integrante 3."),
    (7, "Actores y mapa de relaciones funcionales", 3, 50,
     "Gracias. Identificamos seis actores: transporte, ticketeras, locales gastronómicos, municipio y SERNATUR, seguridad local y los clientes. Sound Travel está al centro porque diseña, contrata, coordina horarios y supervisa en terreno. "
     "Las líneas continuas son relaciones contractuales o comerciales: contratación con el transporte, alianza con las ticketeras, convenio con los locales y compra directa con el cliente. Las punteadas son de coordinación e información: con las instituciones no compramos nada, pero intercambiamos permisos y datos. "
     "Un dato honesto: alojamiento no aplica porque son salidas de un día.",
     ["Sound Travel como coordinador central.", "Dos tipos de vínculo y flujo de información en ambos sentidos."],
     "Veamos qué le corresponde a cada actor."),
    (8, "Responsabilidades y coordinación con cada actor", 3, 40,
     "Cada actor tiene responsabilidades claras: el transporte, máquinas al día y choferes descansados; las ticketeras, entradas válidas y el horario de apertura; los locales, mesas listas a tiempo; y el municipio, SERNATUR y la seguridad, información sobre calles cerradas, estacionamiento de buses y resguardo nocturno. "
     "Pero es una relación de doble vía: ellos también dependen de lo que les entregamos, como nóminas, horarios y pagos en fecha. Abajo está lo comercial: vendemos directo al público por redes, web y WhatsApp. Ojo: con las instituciones es una relación propuesta, no hay convenios firmados.",
     ["Relación de doble vía.", "Vínculo institucional propuesto."],
     "Y todo esto tiene sentido si responde a lo que necesita el cliente."),
    (9, "Necesidades de la demanda e interdependencia", 3, 40,
     "Aquí cruzamos cada necesidad del cliente con nuestra respuesta. Por ejemplo, para volver seguro de noche hay un reencuentro 30 minutos después del show, conteo de pasajeros y una pulsera por bus. "
     "Abajo está la cadena de interdependencia: si la entrada no se emite, el bus no llega o la mesa no está lista, el cliente no queda satisfecho aunque lo demás funcione. Por eso no basta con contratar: hay que coordinar.",
     ["Servicio diseñado desde el cliente.", "Si falla un eslabón, falla el servicio."],
     "¿Con qué criterios aseguramos que cada eslabón cumpla? Le paso la palabra al Integrante 4."),
    (10, "Cuatro criterios para asignar tareas", 4, 50,
     "Gracias. Asignamos las tareas con cuatro criterios. Técnico: antes de cada salida el Jefe de Operaciones revisa neumáticos, cinturones, botiquín y climatización, y el guía usa nómina digital y pulseras. Normativo: proponemos inscribirnos en el registro de SERNATUR y exigimos a los transportistas el Decreto 80, licencia A2 o A3 y SOAP. "
     "Al revisar las leyes descubrimos dos errores de cita en el informe y los corregimos: el registro está en el Título VII de la Ley 20.423 y es voluntario para agencias, y las licencias vienen de la Ley de Tránsito. "
     "En sustentabilidad: Viaje Limpio, cero papel y compras a emprendimientos locales. Y por el territorio, con lluvia, neblina y tacos, usamos puntos de encuentro céntricos e iluminados y salimos con margen.",
     ["Cada criterio define quién hace qué.", "Corregimos dos citas al investigar las normas."],
     "Con estos criterios armamos la matriz de tareas."),
    (11, "Matriz de tareas y seguimiento de la operación", 4, 45,
     "La matriz responde quién hace qué, cuándo y cómo se controla. Antes del evento, el Jefe de Operaciones revisa documentos y confirma a los proveedores por correo siete días antes y por WhatsApp 24 horas antes. El día del evento, guía y conductor hacen el embarque y reportan cuatro hitos en el WhatsApp operativo: bus en el punto, salida, llegada y retorno; si cae la señal, usamos radios. "
     "Al cliente lo acompañamos desde el comprobante hasta la ubicación del guía en tiempo real. Agregamos la fila post-evento, la encuesta, porque en el informe no tenía responsable.",
     ["Seguimiento distinto para proveedores, operadores y clientes.", "Fila post-evento agregada por el equipo."],
     "Y si algo sale mal, tenemos un protocolo."),
    (12, "Contingencias y conclusiones", 4, 45,
     "En verde están las respuestas que ya venían en el informe, como radios, punto de asistencia y pulseras; en ámbar, lo que propusimos para completarlo: proveedores de respaldo y un escalamiento si un pasajero no llega. "
     "Para cerrar: en el criterio 2.1.1 mostramos una estructura y un flujo claros; en el 2.1.2, a Sound Travel coordinando a todos los actores; y en el 2.1.3, tareas asignadas con criterios y controladas con seguimiento. Un servicio de eventos funciona solo si cada actor sabe qué hacer, cuándo y a quién informar. Muchas gracias.",
     ["Distinguir informe y propuesta.", "Un mensaje por criterio."],
     "Quedamos atentos a sus preguntas."),
]

d = base_doc()
t = d.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = t.add_run("SOUND TRAVEL")
r.bold = True
r.font.size = Pt(28)
r.font.color.rgb = VIOLET
r.font.name = "Arial"
p = d.add_paragraph()
r = p.add_run("Guion de exposición para cuatro integrantes")
r.font.size = Pt(16)
r.bold = True
para(d, "Coordinación en Operaciones Turísticas (GTD404) · Unidad 2 · INACAP Sede Temuco · Docente: Evelyn Rios Soto", color=GREY)
para(d, "Versión de 12 diapositivas. Duración total: unos 9 minutos (≈ 2:10 por integrante), con margen para no pasar de 10. Pensado para leerse con naturalidad: es el trabajo que investigamos como equipo.", italic=True)

d.add_heading("1. Distribución de diapositivas y tiempos", level=1)
groups = {1: [], 2: [], 3: [], 4: []}
for n, title, who, secs, *_ in SLIDES:
    groups[who].append((n, title, secs))
rows = []
total = 0
for who in (1, 2, 3, 4):
    secs = sum(s for *_, s in groups[who])
    total += secs
    nums = [n for n, *_ in groups[who]]
    rows.append((f"Integrante {who}", f"{nums[0]}–{nums[-1]}", str(len(nums)),
                 {1: "Portada; problema y Sound Travel (justificación, localización); misión, visión, público y pack",
                  2: "Organigrama y RR. HH.; organigrama funcional y procesos; diagrama de flujo",
                  3: "Actores y mapa de relaciones; responsabilidades y coordinación; demanda e interdependencia",
                  4: "Cuatro criterios; matriz de tareas y seguimiento; contingencias y conclusiones"}[who],
                 {1: "2.1.1 (base)", 2: "2.1.1", 3: "2.1.2", 4: "2.1.3"}[who],
                 f"{secs // 60}:{secs % 60:02d}"))
rows.append(("TOTAL", "1–12", "12", "", "", f"{total // 60}:{total % 60:02d}"))
table(d, ["Expositor", "Diapositivas", "N.°", "Contenido", "Criterio", "Tiempo"], rows, [2.4, 2.4, 1.0, 6.7, 2.0, 1.5])
para(d, "")
para(d, "", bold_lead="Por qué este reparto es equilibrado:")
bullet(d, "Cada integrante expone 3 diapositivas y habla entre 2:10 y 2:20.")
bullet(d, "El Integrante 4 tiene la parte más técnica, así que su guion se queda en las ideas clave; el detalle está en la diapositiva.")
bullet(d, "Se fusionaron diapositivas de la versión larga sin eliminar evidencia: cada subrequisito sigue apareciendo (ver matriz).")
para(d, "Si van justos de tiempo, acorten las diapositivas 2, 7 y 10 (las más largas, 50–60 s). Hay cerca de 1 minuto de margen bajo los 10.", italic=True)

d.add_heading("2. Guion por diapositiva", level=1)
para(d, "Consejo general: no lean la diapositiva. El guion explica y da ejemplos; la diapositiva solo apoya. Las frases entre comillas de transición pueden decirse con sus propias palabras.", italic=True, color=GREY)
current = None
for n, title, who, secs, text, emph, trans in SLIDES:
    if who != current:
        current = who
        d.add_heading(f"INTEGRANTE {who}", level=2)
    d.add_heading(f"Diapositiva {n} · {title}", level=3)
    para(d, f"Integrante {who}   ·   Duración estimada: {secs // 60}:{secs % 60:02d}", color=GREY, size=10)
    para(d, text, bold_lead="Explicación oral: ")
    p = d.add_paragraph()
    r = p.add_run("Ideas que deben enfatizarse:")
    r.bold = True
    for e in emph:
        bullet(d, e)
    para(d, f"«{trans}»", bold_lead="Transición: ", color=VIOLET)

d.add_heading("3. Recomendaciones para el ensayo", level=1)
for tx in [
    "Ensayen con cronómetro al menos dos veces completas, respetando los cambios de integrante.",
    "En cada cambio de expositor, quien termina nombra a quien sigue y el tema que tratará; quien entra comienza con «Gracias».",
    "Si la profesora pregunta por normativa, recuerden: el registro de SERNATUR es voluntario para agencias (Ley 20.423, art. 33) y obligatorio para alojamiento y turismo aventura; las licencias A2/A3 están en la Ley 18.290.",
    "Si preguntan por el festival 2027 o por convenios: respondan que son escenarios y propuestas de una empresa ficticia, no hechos confirmados.",
    "Completen sus nombres en la portada (diapositiva 1) antes de exponer. Usen «investigamos», «revisamos», «propusimos»: el trabajo es de ustedes.",
]:
    bullet(d, tx)
d.save(f"{OUT_DIR}/Sound_Travel_Guion_Exposicion.docx")

# ---------------------------------------------------------------------------
# Matriz de cumplimiento + análisis
# ---------------------------------------------------------------------------
d = base_doc()
p = d.add_paragraph()
r = p.add_run("SOUND TRAVEL · Matriz de cumplimiento académico")
r.bold = True
r.font.size = Pt(20)
r.font.color.rgb = VIOLET
para(d, "GTD404 · Unidad 2 · Relación entre los criterios de la guía, las diapositivas y los expositores", color=GREY)
para(d, "Advertencia: la guía remite a un instrumento de evaluación disponible en AAI que no fue entregado. Esta matriz se construyó con los criterios explícitos de la guía; no es posible garantizar el cumplimiento de indicadores adicionales de esa rúbrica.", italic=True, color=AMBER)

d.add_heading("1. Matriz de verificación", level=1)
M = [
    ("2.1.1", "Selección o creación de la empresa y justificación", "2", "1", "Cumple"),
    ("2.1.1", "Nombre, localización, misión, visión y público objetivo", "2, 3", "1", "Cumple"),
    ("2.1.1", "Propuesta de valor / servicio prestado", "3", "1", "Cumple"),
    ("2.1.1", "Organigrama básico (jerárquico)", "4", "2", "Cumple"),
    ("2.1.1", "Descripción de áreas y funciones", "4", "2", "Cumple"),
    ("2.1.1", "Recurso humano y roles", "4", "2", "Cumple"),
    ("2.1.1", "Organigrama funcional", "5", "2", "Cumple"),
    ("2.1.1", "Procesos internos", "5", "2", "Cumple"),
    ("2.1.1", "Procesos externos: proveedores (alojamiento, transporte, alimentación, seguros)", "5, 7, 8", "2 y 3", "Parcial: transporte, alimentación y ticketeras desarrollados; alojamiento no aplica; seguros solo SOAP"),
    ("2.1.1", "Procesos externos: alianzas estratégicas", "5, 8", "2 y 3", "Cumple (declaradas como propuesta)"),
    ("2.1.1", "Procesos externos: clientes y canales de comercialización", "5, 8", "2 y 3", "Cumple"),
    ("2.1.1", "Diagrama de flujo o mapa de procesos", "6", "2", "Cumple"),
    ("2.1.2", "Identificación de los actores del servicio", "7", "3", "Cumple"),
    ("2.1.2", "Diagrama funcional: empresa al centro, actores alrededor", "7", "3", "Cumple"),
    ("2.1.2", "Líneas según tipo de relación (coordinación, contratación, información)", "7", "3", "Cumple"),
    ("2.1.2", "Responsabilidades y funciones de cada actor", "7, 8", "3", "Cumple"),
    ("2.1.2", "Relación entre servicio y necesidades de la demanda", "3, 9", "1 y 3", "Cumple"),
    ("2.1.2", "Interdependencia operativa entre actores", "7, 8, 9", "3", "Cumple"),
    ("2.1.3", "Criterio de asignación: estándares técnicos", "10", "4", "Cumple"),
    ("2.1.3", "Criterio de asignación: estándares normativos", "10", "4", "Cumple (citas corregidas; verificar texto vigente en bcn.cl)"),
    ("2.1.3", "Criterio de asignación: sustentabilidad", "10", "4", "Cumple"),
    ("2.1.3", "Criterio de asignación: contexto territorial", "2, 10", "1 y 4", "Cumple"),
    ("2.1.3", "Matriz de asignación de tareas y responsables", "11", "4", "Cumple (fila post-evento complementaria)"),
    ("2.1.3", "Seguimiento de proveedores", "11", "4", "Cumple"),
    ("2.1.3", "Seguimiento de operadores", "11", "4", "Cumple"),
    ("2.1.3", "Seguimiento de clientes", "11", "4", "Cumple"),
    ("2.1.3", "Protocolo de contingencias", "12", "4", "Parcial en el informe; completado con propuestas señaladas"),
    ("Forma", "Portada, introducción, desarrollo y conclusión", "1, 2, 3–11, 12", "1–4", "Cumple"),
    ("Forma", "Referencias bibliográficas", "12 (franja) y este documento", "4", "Cumple (versión compacta en la diapositiva)"),
]
table(d, ["Criterio", "Evidencia requerida", "Diapositiva", "Expositor", "Estado"], M, [1.7, 6.4, 2.0, 1.7, 4.2], status_col=4)

d.add_heading("2. Resumen por criterio", level=1)
table(d, ["Criterio", "Evidencia requerida", "Diapositivas", "Expositor", "Estado"], [
    ("2.1.1", "Estructuras y procesos internos/externos", "2–6", "Integrantes 1 y 2", "Cumple (un subrequisito parcial: alojamiento y seguros)"),
    ("2.1.2", "Relaciones funcionales entre actores", "7–9", "Integrante 3", "Cumple"),
    ("2.1.3", "Criterios técnicos, asignación y seguimiento", "10–12", "Integrante 4", "Cumple (contingencias completadas con propuestas)"),
], [1.5, 5.5, 2.2, 3.0, 3.8], status_col=4)

d.add_heading("3. Análisis de los documentos fuente", level=1)
d.add_heading("3.1 Vacíos e inconsistencias detectados", level=2)
gaps = [
    ("Cita normativa (Ley 20.423). ", "El informe atribuye la inscripción en el Registro Nacional de Prestadores al artículo 19. El registro se regula en el Título VII (arts. 30 y siguientes); el art. 33 establece que la inscripción es voluntaria, salvo para alojamiento turístico (art. 34) y turismo aventura. Corregido en la diapositiva 10."),
    ("Cita normativa (licencias). ", "El informe vincula las licencias A2/A3 al D.S. 80/2004. Las clases de licencia están en el art. 12 de la Ley 18.290 de Tránsito (A2: 10 a 17 asientos; A3: sin límite). El D.S. 80 regula el transporte privado remunerado de pasajeros. Corregido en la diapositiva 10."),
    ("Festival Internacional de Temuco 2027. ", "El informe lo presenta como «anunciado». No se encontró confirmación verificable; se presenta como escenario proyectado (diapositiva 2)."),
    ("Proveedores de alojamiento y seguros. ", "La guía los menciona como ejemplos; el informe no los desarrolla. Alojamiento no aplica a salidas de día; para viajes a Santiago o Concepción queda por definir. Seguros: solo SOAP de transportistas; un seguro de asistencia al pasajero no está considerado."),
    ("Matriz de tareas incompleta. ", "El informe cubre cuatro momentos y omite la etapa 5 (encuesta post-evento), sin responsable asignado. Se agregó una fila complementaria (diapositiva 11)."),
    ("Protocolo de contingencias. ", "El informe menciona planes de contingencia y medidas aisladas, pero no un protocolo formal. Se ordenaron en una tabla y se añadieron dos situaciones y un escalamiento marcados como propuesta (diapositiva 12)."),
    ("Responsables de etapas 1 y 2 del flujo. ", "No están explícitos en la figura 2; se dedujeron de la secc. 3.1 y 3.2 (área Comercial)."),
    ("Cantidad de integrantes. ", "La guía indica grupos de 2 o 3 integrantes y el informe registra un solo integrante; la presentación se diseñó para 4 según lo solicitado. Confirmen con la docente si la conformación es válida."),
    ("Formato de entrega. ", "La guía describe un pre-informe escrito y su entrega como «material audiovisual» en AAI; la presentación es una adaptación oral del informe."),
    ("Instrumento de evaluación. ", "La rúbrica de AAI no fue proporcionada (ver advertencia inicial)."),
]
for lead, tx in gaps:
    bullet(d, tx, lead)

d.add_heading("3.2 Clasificación de la información utilizada", level=2)
table(d, ["Tipo", "Ejemplos en la presentación", "Cómo se presenta"], [
    ("Dato verificable externo", "3,6 millones de visitantes internacionales en 2023 (Subsecretaría de Turismo, 2024, citada en la guía); contenido de Ley 20.423, D.S. 80, Ley 18.290 y Ley 18.490", "Con fuente en la diapositiva"),
    ("Propuesta operativa de la empresa ficticia", "Estructura, cargos, rutas, pack, convenios, alianzas institucionales, matriz, canales de seguimiento, sustentabilidad", "Atribuida al informe; convenios e inscripción en registro rotulados como propuesta"),
    ("Escenario proyectado", "Festival Internacional de Temuco 2027", "Etiqueta «Escenario proyectado»"),
    ("Complemento del equipo", "Fila post-evento, dos contingencias, escalamiento, cadena de interdependencia", "Rotulado en ámbar o como «análisis del equipo»"),
    ("Requiere verificación", "Texto vigente de las normas en bcn.cl; condiciones reales de tránsito y permisos municipales para buses turísticos", "Nota en diapositivas 10 y 12"),
], [3.5, 8.0, 4.5])

d.add_heading("3.3 Fuentes consultadas para la verificación normativa", level=2)
for tx in [
    "SERNATUR – Ley 20.423 modificada: https://www.sernatur.cl/transparencia/archivos/marco-normativo/marco-norm-aplic/LEY-20423-modificada.pdf",
    "SERNATUR – Portal de Servicios Turísticos (registro): https://portalserviciosturisticos.sernatur.cl/formalizate/sello-r/",
    "BCN – Ley 20.423: https://www.bcn.cl/leychile/navegar?i=1010960",
    "BCN – D.S. 80/2004 MTT: https://www.bcn.cl/leychile/navegar?i=230180",
    "SUSESO – Ley 18.290, art. 12: https://www.suseso.gob.cl/612/w3-propertyvalue-99407.html",
]:
    bullet(d, tx)
para(d, "Recomendación: antes de la exposición, revisen el texto vigente de cada norma directamente en bcn.cl.", italic=True, color=AMBER)
d.save(f"{OUT_DIR}/Sound_Travel_Matriz_Cumplimiento.docx")
print("ok")
